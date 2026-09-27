"""
Unit tests for multi-sheet Excel dossier export engine.
"""
import io
import openpyxl
from src.data_exporter import generate_excel_crop_dossier


def test_excel_dossier_generation_default():
    soil = {"nitrogen": 90.0, "phosphorus": 42.0, "potassium": 43.0, "ph": 6.5}
    climate = {"temperature": 26.5, "humidity": 75.0, "rainfall": 110.0}
    
    excel_bytes = generate_excel_crop_dossier(
        crop_name="Rice",
        viability=0.95,
        soil_profile=soil,
        climate_profile=climate,
        field_area_acres=2.5
    )
    
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000
    
    # Load with openpyxl to verify sheet names and content
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Executive Summary" in sheet_names
    assert "Nutrient Prescription" in sheet_names
    assert "Economic Feasibility" in sheet_names
    
    ws_summary = wb["Executive Summary"]
    assert ws_summary["B5"].value == "Rice"
    assert ws_summary["B7"].value == 2.5
    
    ws_fert = wb["Nutrient Prescription"]
    assert ws_fert.max_row >= 4


def test_excel_dossier_with_custom_economics():
    soil = {"nitrogen": 80.0, "phosphorus": 35.0, "potassium": 40.0, "ph": 6.8}
    climate = {"temperature": 28.0, "humidity": 65.0, "rainfall": 90.0}
    
    custom_econ = {
        "financial_summary": {
            "gross_revenue_inr": 85000.0,
            "total_cost_inr": 32000.0,
            "net_profit_inr": 53000.0,
            "benefit_cost_ratio": 2.65
        }
    }
    
    excel_bytes = generate_excel_crop_dossier(
        crop_name="Cotton",
        viability=0.88,
        soil_profile=soil,
        climate_profile=climate,
        field_area_acres=1.0,
        economics_data=custom_econ
    )
    
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws_econ = wb["Economic Feasibility"]
    assert ws_econ["C4"].value == 85000.0
    assert ws_econ["C5"].value == 32000.0
    assert ws_econ["C6"].value == 53000.0
